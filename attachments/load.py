"""Loader functions that transform files into attachment objects."""

from . import matchers
from .core import Attachment, loader, AttachmentCollection
import io

# --- URL DOWNLOADER ---
@loader(match=matchers.url_match)
def url_to_bs4(att: Attachment) -> Attachment:
    """Load URL content and parse with BeautifulSoup."""
    try:
        import requests
        from bs4 import BeautifulSoup
        
        response = requests.get(att.path, timeout=10)
        response.raise_for_status()
        
        # Parse with BeautifulSoup
        soup = BeautifulSoup(response.content, 'html.parser')
        
        # Store the soup object
        att._obj = soup
        # Store some metadata
        att.metadata.update({
            'content_type': response.headers.get('content-type', ''),
            'status_code': response.status_code,
        })
        
        return att
    except ImportError:
        raise ImportError("requests and beautifulsoup4 are required for URL loading. Install with: pip install requests beautifulsoup4")

@loader(match=lambda att: att.path.startswith(('http://', 'https://')) and any(att.path.lower().endswith(ext) for ext in ['.pdf', '.pptx', '.ppt', '.docx', '.doc', '.xlsx', '.xls', '.csv', '.jpg', '.jpeg', '.png', '.gif', '.bmp']))
def url_to_file(att: Attachment) -> Attachment:
    """Download file from URL and delegate to appropriate loader based on file extension."""
    try:
        import requests
        import tempfile
        import os
        from urllib.parse import urlparse
        from pathlib import Path
        
        # Parse URL to get file extension
        parsed_url = urlparse(att.path)
        url_path = parsed_url.path
        
        # Get file extension from URL
        file_ext = Path(url_path).suffix.lower()
        
        # Download the file
        response = requests.get(att.path, timeout=30)
        response.raise_for_status()
        
        # Create temporary file with correct extension
        with tempfile.NamedTemporaryFile(suffix=file_ext, delete=False) as temp_file:
            temp_file.write(response.content)
            temp_path = temp_file.name
        
        # Store original URL and temp path
        original_url = att.path
        att.path = temp_path
        att.metadata.update({
            'original_url': original_url,
            'temp_file_path': temp_path,
            'downloaded_from_url': True,
            'content_length': len(response.content),
            'content_type': response.headers.get('content-type', ''),
        })
        
        # Now delegate to the appropriate loader based on file extension
        if file_ext in ('.pdf',):
            return pdf_to_pdfplumber(att)
        elif file_ext in ('.pptx', '.ppt'):
            return pptx_to_python_pptx(att)
        elif file_ext in ('.docx', '.doc'):
            return docx_to_python_docx(att)
        elif file_ext in ('.xlsx', '.xls'):
            return excel_to_openpyxl(att)
        elif file_ext in ('.csv',):
            return csv_to_pandas(att)
        elif file_ext.lower() in ('.jpg', '.jpeg', '.png', '.gif', '.bmp', '.heic', '.heif'):
            return image_to_pil(att)
        else:
            # If we don't recognize the extension, try to guess from content-type
            content_type = response.headers.get('content-type', '').lower()
            if 'pdf' in content_type:
                return pdf_to_pdfplumber(att)
            elif 'powerpoint' in content_type or 'presentation' in content_type:
                return pptx_to_python_pptx(att)
            elif 'word' in content_type or 'document' in content_type:
                return docx_to_python_docx(att)
            elif 'excel' in content_type or 'spreadsheet' in content_type:
                return excel_to_openpyxl(att)
            else:
                # Fallback: treat as text
                att._obj = response.text
                att.text = response.text
                return att
                
    except ImportError:
        raise ImportError("requests is required for URL loading. Install with: pip install requests")
    except Exception as e:
        # Clean up temp file if it was created
        if 'temp_file_path' in att.metadata:
            try:
                os.unlink(att.metadata['temp_file_path'])
            except:
                pass
        raise ValueError(f"Could not download or process URL: {e}")

# --- LOADERS ---
@loader(match=matchers.csv_match)
def csv_to_pandas(att: Attachment) -> Attachment:
    """Load CSV into pandas DataFrame."""
    try:
        import pandas as pd
        att._obj = pd.read_csv(att.path)
    except ImportError:
        raise ImportError("pandas is required for CSV loading. Install with: pip install pandas")
    return att


@loader(match=matchers.pdf_match)
def pdf_to_pdfplumber(att: Attachment) -> Attachment:
    """Load PDF using pdfplumber."""
    try:
        import pdfplumber
        
        # Try to create a temporary PDF with CropBox defined to silence warnings
        try:
            import pypdf
            from io import BytesIO
            import tempfile
            import os
            
            # Read the original PDF
            with open(att.path, 'rb') as f:
                pdf_bytes = f.read()
            
            # Process with pypdf to add CropBox
            reader = pypdf.PdfReader(BytesIO(pdf_bytes))
            writer = pypdf.PdfWriter()
            
            for page in reader.pages:
                # Set CropBox to MediaBox if not already defined
                if '/CropBox' not in page:
                    page.cropbox = page.mediabox
                writer.add_page(page)
            
            # Create a temporary file with the modified PDF
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as temp_file:
                writer.write(temp_file)
                temp_path = temp_file.name
            
            # Open the temporary PDF with pdfplumber
            att._obj = pdfplumber.open(temp_path)
            
            # Store the temp path for cleanup later
            att.metadata['temp_pdf_path'] = temp_path
            
        except (ImportError, Exception):
            # If CropBox fix fails, fall back to original file
            att._obj = pdfplumber.open(att.path)
            
    except ImportError:
        raise ImportError("pdfplumber is required for PDF loading. Install with: pip install pdfplumber")
    return att


@loader(match=matchers.pptx_match)
def pptx_to_python_pptx(att: Attachment) -> Attachment:
    """Load PowerPoint using python-pptx."""
    try:
        from pptx import Presentation
        att._obj = Presentation(att.path)
    except ImportError:
        raise ImportError("python-pptx is required for PowerPoint loading. Install with: pip install python-pptx")
    return att


@loader(match=matchers.docx_match)
def docx_to_python_docx(att: Attachment) -> Attachment:
    """Load Word document using python-docx."""
    try:
        from docx import Document
        att._obj = Document(att.path)
    except ImportError:
        raise ImportError("python-docx is required for Word document loading. Install with: pip install python-docx")
    return att


@loader(match=matchers.excel_match)
def excel_to_openpyxl(att: Attachment) -> Attachment:
    """Load Excel workbook using openpyxl."""
    try:
        from openpyxl import load_workbook
        att._obj = load_workbook(att.path, read_only=True)
    except ImportError:
        raise ImportError("openpyxl is required for Excel loading. Install with: pip install openpyxl")
    return att


@loader(match=matchers.image_match)
def image_to_pil(att: Attachment) -> Attachment:
    """Load image using PIL."""
    try:
        # Try to import pillow-heif for HEIC support if needed
        if att.path.lower().endswith(('.heic', '.heif')):
            try:
                from pillow_heif import register_heif_opener
                register_heif_opener()
            except ImportError:
                pass  # Fall back to PIL's built-in support if available
        
        from PIL import Image
        att._obj = Image.open(att.path)
        
        # Store metadata
        att.metadata.update({
            'format': getattr(att._obj, 'format', 'Unknown'),
            'size': getattr(att._obj, 'size', (0, 0)),
            'mode': getattr(att._obj, 'mode', 'Unknown')
        })
        
    except ImportError:
        if att.path.lower().endswith(('.heic', '.heif')):
            raise ImportError("pillow-heif is required for HEIC loading. Install with: pip install pillow-heif")
        else:
            raise ImportError("Pillow is required for image loading. Install with: pip install Pillow")
    return att


@loader(match=matchers.text_match)
def text_to_string(att: Attachment) -> Attachment:
    """Load text files as strings."""
    with open(att.path, 'r', encoding='utf-8') as f:
        content = f.read()
        att._obj = content
        att.text = content
    return att


@loader(match=lambda att: att.path.lower().endswith(('.html', '.htm')))
def html_to_bs4(att: Attachment) -> Attachment:
    """Load HTML files and parse with BeautifulSoup."""
    try:
        from bs4 import BeautifulSoup
        
        with open(att.path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Parse with BeautifulSoup
        soup = BeautifulSoup(content, 'html.parser')
        
        # Store the soup object
        att._obj = soup
        # Store some metadata
        att.metadata.update({
            'content_type': 'text/html',
            'file_size': len(content),
        })
        
        return att
    except ImportError:
        raise ImportError("beautifulsoup4 is required for HTML loading. Install with: pip install beautifulsoup4")


@loader(match=matchers.zip_match)
def zip_to_images(att: Attachment) -> 'AttachmentCollection':
    """Load ZIP file containing images into AttachmentCollection."""
    try:
        import zipfile
        from PIL import Image
        from .core import AttachmentCollection, Attachment
        
        attachments = []
        
        with zipfile.ZipFile(att.path, 'r') as zip_file:
            for file_info in zip_file.filelist:
                if file_info.filename.lower().endswith(('.jpg', '.jpeg', '.png', '.gif', '.bmp', '.heic', '.heif')):
                    # Create attachment for each image
                    img_att = Attachment(file_info.filename)
                    
                    # Copy commands from original attachment (for vectorized processing)
                    img_att.commands = att.commands.copy()
                    
                    # Load image from zip
                    with zip_file.open(file_info.filename) as img_file:
                        img_data = img_file.read()
                        img = Image.open(io.BytesIO(img_data))
                        img_att._obj = img
                        
                        # Store metadata
                        img_att.metadata.update({
                            'format': getattr(img, 'format', 'Unknown'),
                            'size': getattr(img, 'size', (0, 0)),
                            'mode': getattr(img, 'mode', 'Unknown'),
                            'from_zip': att.path,
                            'zip_filename': file_info.filename
                        })
                    
                    attachments.append(img_att)
        
        return AttachmentCollection(attachments)
        
    except ImportError:
        raise ImportError("Pillow is required for image processing. Install with: pip install Pillow")
    except Exception as e:
        raise ValueError(f"Could not load ZIP file: {e}")
